"""Documents built directly (no fragile UI typing): PowerPoint decks from an outline."""
import re
from pathlib import Path
from typing import List, Tuple


def parse_outline(outline: str) -> List[Tuple[str, List[str]]]:
    slides, cur = [], None
    for raw in outline.replace("\r\n", "\n").split("\n"):
        line = raw.strip()
        if not line:
            continue
        m = re.match(r"^(#{1,3}|slide\s*\d*\s*[:.-])\s*(.+)$", line, re.I)
        if m:
            cur = (m.group(2).strip(), [])
            slides.append(cur)
        elif cur is not None:
            cur[1].append(re.sub(r"^[-*\u2022\d.)\s]+", "", line).strip())
        else:
            cur = (line, [])
            slides.append(cur)
    return [s for s in slides if s[0]]


def build_pptx(title: str, outline: str, path: Path) -> Path:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt
    slides = parse_outline(outline) or [(title, [])]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    ink, accent, muted = RGBColor(0x14, 0x18, 0x20), RGBColor(0x7C, 0x5C, 0xFF), RGBColor(0x5A, 0x64, 0x78)

    def bar(s):
        shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), prs.slide_height)
        shp.fill.solid(); shp.fill.fore_color.rgb = accent; shp.line.fill.background()

    first_title, first_body = slides[0]
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar(s)
    tb = s.shapes.add_textbox(Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.6)).text_frame
    tb.word_wrap = True
    tb.text = first_title
    tb.paragraphs[0].runs[0].font.size, tb.paragraphs[0].runs[0].font.bold = Pt(48), True
    tb.paragraphs[0].runs[0].font.color.rgb = ink
    if first_body:
        p = tb.add_paragraph(); p.text = first_body[0]
        p.runs[0].font.size, p.runs[0].font.color.rgb = Pt(22), muted
    for head, bullets in slides[1:]:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        bar(s)
        h = s.shapes.add_textbox(Inches(0.9), Inches(0.6), Inches(11.5), Inches(1.1)).text_frame
        h.word_wrap = True
        h.text = head
        h.paragraphs[0].runs[0].font.size, h.paragraphs[0].runs[0].font.bold = Pt(34), True
        h.paragraphs[0].runs[0].font.color.rgb = ink
        body = s.shapes.add_textbox(Inches(0.9), Inches(1.9), Inches(11.5), Inches(5)).text_frame
        body.word_wrap = True
        for i, b in enumerate(bullets[:8]):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            p.text = "\u2022  " + b
            p.space_after = Pt(12)
            p.runs[0].font.size, p.runs[0].font.color.rgb = Pt(22), ink
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path
